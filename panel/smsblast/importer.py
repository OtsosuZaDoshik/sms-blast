"""Импорт базы номеров из CSV и Excel."""

import csv
import io
import os
import re

from . import db, phones

PHONE_HEADERS = ["телефон", "тел", "номер", "phone", "mobile", "сотовый", "моб", "msisdn"]
NAME_HEADERS = ["имя", "name", "фио", "клиент", "контакт", "получатель", "фамилия"]
CONSENT_HEADERS = ["согласие", "consent", "opt-in", "optin", "подписка"]

TRUTHY = {"1", "да", "true", "yes", "y", "истина", "+", "есть", "согласен", "согласна"}


def _decode(raw):
    """CSV из Excel на macOS часто приходит в cp1251 или utf-8 с BOM."""
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "koi8-r"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_table(path):
    """Читает CSV/XLSX → (заголовки, список словарей)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path)
    return _read_csv(path)


def _read_csv(path):
    with open(path, "rb") as handle:
        text = _decode(handle.read())

    sample = text[:8192]
    first_line = sample.splitlines()[0] if sample.strip() else ""

    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = None

    # Sniffer ошибается, когда внутри значений есть символ другого разделителя
    # (например JSON с запятыми в файле с «;»). Проверяем результат: настоящий
    # разделитель обязан разрезать шапку больше чем на одну колонку.
    if not delimiter or len(next(csv.reader([first_line], delimiter=delimiter))) < 2:
        counts = {d: len(next(csv.reader([first_line], delimiter=d)))
                  for d in (";", ",", "\t", "|")}
        best = max(counts, key=lambda d: counts[d])
        delimiter = best if counts[best] > 1 else ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [r for r in reader if any(str(c).strip() for c in r)]
    if not rows:
        return [], []

    headers = [str(h).strip() for h in rows[0]]
    if not _looks_like_header(headers):
        headers = ["колонка {}".format(i + 1) for i in range(len(rows[0]))]
        body = rows
    else:
        body = rows[1:]

    return headers, [_zip_row(headers, r) for r in body]


def _read_excel(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if row and any(cell is not None and str(cell).strip() for cell in row):
            rows.append(["" if c is None else str(c).strip() for c in row])
    workbook.close()
    if not rows:
        return [], []

    headers = rows[0]
    if not _looks_like_header(headers):
        headers = ["колонка {}".format(i + 1) for i in range(len(rows[0]))]
        body = rows
    else:
        body = rows[1:]
    return headers, [_zip_row(headers, r) for r in body]


def _looks_like_header(cells):
    """Шапка есть, если первая строка не состоит из телефонов."""
    joined = " ".join(str(c).lower() for c in cells)
    if any(h in joined for h in PHONE_HEADERS + NAME_HEADERS):
        return True
    digit_cells = sum(1 for c in cells if re.sub(r"\D", "", str(c)) and
                      len(re.sub(r"\D", "", str(c))) >= 10)
    return digit_cells == 0


def _zip_row(headers, row):
    data = {}
    for index, header in enumerate(headers):
        data[header] = str(row[index]).strip() if index < len(row) else ""
    return data


def guess_column(headers, candidates):
    lowered = [str(h).strip().lower() for h in headers]
    for index, header in enumerate(lowered):
        if header in candidates:
            return headers[index]
    for index, header in enumerate(lowered):
        if any(c in header for c in candidates):
            return headers[index]
    return None


def guess_phone_column(headers, rows):
    column = guess_column(headers, PHONE_HEADERS)
    if column:
        return column
    # Шапка не помогла — берём колонку, где больше всего валидных номеров.
    best, best_hits = None, 0
    for header in headers:
        hits = 0
        for row in rows[:50]:
            number, error = phones.normalize(row.get(header))
            if number and not error:
                hits += 1
        if hits > best_hits:
            best, best_hits = header, hits
    return best if best_hits else None


def import_rows(rows, phone_column, name_column=None, consent_column=None,
                default_consent=True, country_code="7", source=""):
    """Пишет контакты в БД. Возвращает сводку импорта."""
    result = {"added": 0, "updated": 0, "invalid": 0, "duplicates": 0,
              "optout": 0, "errors": []}
    seen = set()

    for index, row in enumerate(rows, start=2):
        phone, error = phones.normalize(row.get(phone_column), country_code)
        if error:
            result["invalid"] += 1
            if len(result["errors"]) < 20:
                result["errors"].append(
                    "строка {}: «{}» — {}".format(index, row.get(phone_column, ""), error)
                )
            continue

        if phone in seen:
            result["duplicates"] += 1
            continue
        seen.add(phone)

        if db.is_optout(phone):
            result["optout"] += 1
            continue

        consent = 1 if default_consent else 0
        if consent_column:
            raw = str(row.get(consent_column, "")).strip().lower()
            consent = 1 if raw in TRUTHY else 0

        name = str(row.get(name_column, "")).strip() if name_column else ""
        fields = {k: v for k, v in row.items() if k and k != phone_column}

        action = db.upsert_contact(phone, name, fields, consent, source)
        result[action] += 1

    return result
